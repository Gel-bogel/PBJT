import streamlit as st
import pandas as pd
import numpy as np
import io
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="Laporan Realisasi PPJ - UP3 Medan Utara", layout="wide")

# CSS Styling untuk Web Dashboard
st.markdown("""
<style>
    .report-table {
        width: 100%;
        border-collapse: collapse;
        font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
        font-size: 13px;
        color: #000000;
        margin-bottom: 25px;
        table-layout: fixed;
    }
    .report-table th, .report-table td {
        border: 1px solid #A6ACAF;
        padding: 6px 8px;
    }
    .col-no {
        width: 45px;
        min-width: 45px;
        max-width: 45px;
        text-align: center;
    }
    .col-ulp {
        width: 170px;
        min-width: 170px;
        max-width: 170px;
        text-align: left;
    }
    .col-data {
        width: 175px;
        min-width: 175px;
        max-width: 175px;
    }
    .curr-cell {
        display: flex;
        justify-content: space-between;
        align-items: center;
        width: 100%;
        font-variant-numeric: tabular-nums;
    }
    .curr-symbol {
        font-weight: 500;
        color: #333333;
        margin-right: 6px;
    }
</style>
""", unsafe_allow_html=True)

st.title("⚡ Laporan Realisasi Hasil Pemungutan PPJ / PBJT-TL")
st.subheader("PT PLN (PERSERO) UP3 MEDAN UTARA")
st.caption("Konsolidasi 5 ULP (Format Baru AP2T)")

# Sidebar Upload File
st.sidebar.header("📁 Unggah Berkas AP2T (Format Baru)")
uploaded_ppj1_files = st.sidebar.file_uploader(
    "1. Berkas Halaman 1 (PPJ 1)", 
    type=["xls", "xlsx"], 
    accept_multiple_files=True
)
uploaded_ppj2_files = st.sidebar.file_uploader(
    "2. Berkas Halaman 2 (PPJ 2)", 
    type=["xls", "xlsx"], 
    accept_multiple_files=True
)

def get_exact_ulp_name(df):
    """Mendeteksi nama ULP secara presisi dari header file AP2T"""
    for r in range(min(7, len(df))):
        for c in range(min(10, df.shape[1])):
            val = str(df.iloc[r, c]).strip().upper()
            if "ULP " in val:
                parts = val.split("ULP ")
                if len(parts) > 1:
                    return "ULP " + parts[1].split()[0]
    return "ULP LAINNYA"

def extract_ppj1_pemda_new(df, pemda_name):
    """Ekstraksi PBJT-TL Kode Gerak Format Baru (Halaman 1)"""
    mask = (df.iloc[:, 2].astype(str).str.strip().str.upper() == pemda_name.upper()) & \
           (df.iloc[:, 5].astype(str).str.strip().str.upper().str.startswith("JUMLAH"))
    df_sub = df[mask]
    
    if df_sub.empty:
        mask = (df.iloc[:, 2].astype(str).str.strip().str.upper() == pemda_name.upper()) & \
               (pd.to_numeric(df.iloc[:, 31], errors='coerce').notnull())
        df_sub = df[mask]

    # Berdasarkan struktur format baru Halaman 1:
    # Kolom PBJT-TL/PPJ per Kode Gerak
    k4 = pd.to_numeric(df_sub.iloc[:, 14], errors='coerce').sum()  # Kode Gerak 21 (Lunas Offline)
    k3 = pd.to_numeric(df_sub.iloc[:, 18], errors='coerce').sum()  # Kode Gerak 22 (Lunas Online)
    # Kode Gerak 23 diabaikan
    k5 = pd.to_numeric(df_sub.iloc[:, 26], errors='coerce').sum()  # Kode Gerak 24 (Beban Kantor)
    k6 = pd.to_numeric(df_sub.iloc[:, 30], errors='coerce').sum()  # Kode Gerak 25 (Lunas Terpusat)

    return {"k3": k3, "k4": k4, "k5": k5, "k6": k6}

def extract_ppj2_pemda_new(df, pemda_name):
    """Ekstraksi PBJT-TL Format Baru (Halaman 2)"""
    mask = (df.iloc[:, 1].astype(str).str.strip().str.upper() == pemda_name.upper()) & \
           (df.iloc[:, 3].astype(str).str.strip().str.upper().str.startswith("JUMLAH"))
    df_sub = df[mask]
    
    if df_sub.empty:
        mask = (df.iloc[:, 1].astype(str).str.strip().str.upper() == pemda_name.upper()) & \
               (pd.to_numeric(df.iloc[:, 48], errors='coerce').notnull())
        df_sub = df[mask]

    # Restitusi (Taglis, Non-Taglis, Prabayar)
    res_taglis = pd.to_numeric(df_sub.iloc[:, 11], errors='coerce').sum()
    res_nontaglis = pd.to_numeric(df_sub.iloc[:, 15], errors='coerce').sum()
    res_prabayar = pd.to_numeric(df_sub.iloc[:, 19], errors='coerce').sum()
    k7 = res_taglis + res_nontaglis + res_prabayar

    # Non-Taglis (Online, Offline, Nota Buku)
    k8 = pd.to_numeric(df_sub.iloc[:, 23], errors='coerce').sum()
    k9 = pd.to_numeric(df_sub.iloc[:, 27], errors='coerce').sum()
    k10 = pd.to_numeric(df_sub.iloc[:, 31], errors='coerce').sum()

    # Prabayar (Token Online, TMP, UJL)
    k11 = pd.to_numeric(df_sub.iloc[:, 35], errors='coerce').sum()
    k12 = pd.to_numeric(df_sub.iloc[:, 40], errors='coerce').sum()
    k13 = pd.to_numeric(df_sub.iloc[:, 44], errors='coerce').sum()

    return {
        "k7": k7, "k8": k8, "k9": k9, "k10": k10,
        "k11": k11, "k12": k12, "k13": k13
    }

# Mapping berkas berdasarkan ULP
dict_ppj1 = {}
dict_ppj2 = {}

if uploaded_ppj1_files:
    for f in uploaded_ppj1_files:
        df = pd.read_excel(f, header=None)
        ulp = get_exact_ulp_name(df)
        dict_ppj1[ulp] = df

if uploaded_ppj2_files:
    for f in uploaded_ppj2_files:
        df = pd.read_excel(f, header=None)
        ulp = get_exact_ulp_name(df)
        dict_ppj2[ulp] = df

if dict_ppj1 and dict_ppj2:
    all_ulps = sorted(list(set(dict_ppj1.keys()).union(set(dict_ppj2.keys()))))
    
    rows_medan = []
    rows_deli = []

    for ulp in all_ulps:
        df1 = dict_ppj1.get(ulp)
        df2 = dict_ppj2.get(ulp)

        d1_m = extract_ppj1_pemda_new(df1, "KOTA MEDAN") if df1 is not None else {"k3": 0, "k4": 0, "k5": 0, "k6": 0}
        d2_m = extract_ppj2_pemda_new(df2, "KOTA MEDAN") if df2 is not None else {"k7": 0, "k8": 0, "k9": 0, "k10": 0, "k11": 0, "k12": 0, "k13": 0}
        row_m = {"ULP": ulp}
        row_m.update(d1_m)
        row_m.update(d2_m)
        rows_medan.append(row_m)

        d1_ds = extract_ppj1_pemda_new(df1, "KAB. DELI SERDANG") if df1 is not None else {"k3": 0, "k4": 0, "k5": 0, "k6": 0}
        d2_ds = extract_ppj2_pemda_new(df2, "KAB. DELI SERDANG") if df2 is not None else {"k7": 0, "k8": 0, "k9": 0, "k10": 0, "k11": 0, "k12": 0, "k13": 0}
        row_ds = {"ULP": ulp}
        row_ds.update(d1_ds)
        row_ds.update(d2_ds)
        rows_deli.append(row_ds)

    def build_table(rows):
        df = pd.DataFrame(rows)
        # 15 = (3 + 4 + 5 + 6 - 7 + 8 + 9 + 10 + 11 - 12 + 13)
        df["k15"] = (
            df["k3"] + df["k4"] + df["k5"] + df["k6"]
            - df["k7"]
            + df["k8"] + df["k9"] + df["k10"] + df["k11"]
            - df["k12"]
            + df["k13"]
        )
        df.insert(0, "NO", range(1, len(df) + 1))
        
        num_cols = [f"k{i}" for i in [3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 15]]
        tot_vals = df[num_cols].sum()
        tot_dict = {"NO": "", "ULP": "JUMLAH TOTAL UP3 MEDAN UTARA"}
        tot_dict.update(tot_vals.to_dict())
        df_full = pd.concat([df, pd.DataFrame([tot_dict])], ignore_index=True)
        return df_full

    df_medan_raw = build_table(rows_medan)
    df_deli_raw = build_table(rows_deli)

    # Tabel 1 Keseluruhan
    num_cols = [f"k{i}" for i in [3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 15]]
    df_total_all = df_medan_raw.copy()
    for col in num_cols:
        df_total_all[col] = df_medan_raw[col] + df_deli_raw[col]
    df_total_all["ULP"] = df_medan_raw["ULP"]
    df_total_all["NO"] = df_medan_raw["NO"]

    col_mapping = {
        "NO": "NO",
        "ULP": "NAMA ULP",
        "k3": "LUNAS ONLINE (KODE GERAK 22)",
        "k4": "LUNAS OFFLINE (KODE GERAK 21)",
        "k5": "BEBAN KANTOR (KODE GERAK 24)",
        "k6": "LUNAS TERPUSAT (KODE GERAK 25)",
        "k7": "RESTITUSI",
        "k8": "NON TAGLIS ONLINE",
        "k9": "NON TAGLIS OFFLINE",
        "k10": "NOTA BUKU",
        "k11": "PRABAYAR",
        "k12": "Kompensasi TMP",
        "k13": "Kompensasi UJL",
        "k15": "JUMLAH [15=(3+4+5+6-7+8+9+10+11-12+13)]"
    }

    def render_tidy_table(df_data):
        df_disp = df_data.rename(columns=col_mapping)
        col_numbers = ["1", "2", "3", "4", "5", "6", "7", "8", "9", "10", "11", "12", "13", "15"]
        
        html = '<div style="overflow-x:auto;"><table class="report-table">'
        html += '<thead><tr style="background-color: #BCE5F7; text-align: center; font-weight: bold;">'
        for idx, num in enumerate(col_numbers):
            cls = "col-no" if idx == 0 else ("col-ulp" if idx == 1 else "col-data")
            html += f'<th class="{cls}" style="border: 1px solid #7EC8E3; padding: 5px; color: #000000;">{num}</th>'
        html += '</tr>'
        
        html += '<tr style="background-color: #7EC8E3; text-align: center; font-weight: bold; color: #000000;">'
        for idx, col_name in enumerate(df_disp.columns):
            cls = "col-no" if idx == 0 else ("col-ulp" if idx == 1 else "col-data")
            html += f'<th class="{cls}" style="border: 1px solid #5DADE2; padding: 8px 6px; vertical-align: middle;">{col_name}</th>'
        html += '</tr></thead><tbody>'
        
        for _, row in df_disp.iterrows():
            is_total = "JUMLAH TOTAL UP3" in str(row["NAMA ULP"])
            if is_total:
                row_style = 'background-color: #7EC8E3; font-weight: bold; border-top: 2px solid #2980B9; border-bottom: 3px double #000000;'
            else:
                row_style = 'background-color: #FFFFFF;'
                
            html += f'<tr style="{row_style}">'
            for idx, val in enumerate(row):
                if idx == 0:
                    html += f'<td class="col-no">{val}</td>'
                elif idx == 1:
                    html += f'<td class="col-ulp" style="white-space: nowrap;">{val}</td>'
                else:
                    num_val = float(val) if pd.notnull(val) else 0.0
                    formatted_num = f"{num_val:,.0f}"
                    html += f'<td class="col-data"><div class="curr-cell"><span class="curr-symbol">Rp</span><span>{formatted_num}</span></div></td>'
            html += '</tr>'
        html += '</tbody></table></div>'
        st.markdown(html, unsafe_allow_html=True)

    # Tampilan Dashboard
    st.markdown("### 🏢 TABEL 1. TOTAL KESELURUHAN (PEMDA MEDAN + DELI SERDANG)")
    render_tidy_table(df_total_all)
    st.markdown("---")
    st.markdown("### 🏛️ TABEL 2. PEMDA KOTA MEDAN")
    render_tidy_table(df_medan_raw)
    st.markdown("---")
    st.markdown("### 🌾 TABEL 3. PEMDA KAB. DELI SERDANG")
    render_tidy_table(df_deli_raw)

    # Ekspor Excel
    def style_excel_sheet(ws, title_text, col_numbers):
        ws.merge_cells("A1:N1")
        ws["A1"] = title_text
        ws["A1"].font = Font(name="Arial", size=13, bold=True, color="000000")
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

        header_num_fill = PatternFill(start_color="BCE5F7", end_color="BCE5F7", fill_type="solid")
        font_dark_bold = Font(name="Arial", size=10, bold=True, color="000000")

        for c_idx, num in enumerate(col_numbers, start=1):
            cell = ws.cell(row=3, column=c_idx, value=num)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.font = font_dark_bold
            cell.fill = header_num_fill

        header_name_fill = PatternFill(start_color="7EC8E3", end_color="7EC8E3", fill_type="solid")
        thin_border = Border(left=Side(style='thin', color='A6ACAF'), right=Side(style='thin', color='A6ACAF'), top=Side(style='thin', color='A6ACAF'), bottom=Side(style='thin', color='A6ACAF'))

        for c_idx in range(1, len(col_numbers) + 1):
            cell = ws.cell(row=4, column=c_idx)
            cell.fill = header_name_fill
            cell.font = font_dark_bold
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

        total_fill = PatternFill(start_color="7EC8E3", end_color="7EC8E3", fill_type="solid")
        double_bottom_border = Border(top=Side(style='thin', color='2980B9'), bottom=Side(style='double', color='000000'), left=Side(style='thin', color='A6ACAF'), right=Side(style='thin', color='A6ACAF'))

        max_r = ws.max_row
        for r in range(5, max_r + 1):
            is_total_row = (r == max_r)
            for c in range(1, len(col_numbers) + 1):
                cell = ws.cell(row=r, column=c)
                if is_total_row:
                    cell.fill = total_fill
                    cell.font = font_dark_bold
                    cell.border = double_bottom_border
                else:
                    cell.border = thin_border
                    cell.font = Font(name="Arial", size=10, color="000000")
                if c >= 3 and isinstance(cell.value, (int, float)):
                    cell.number_format = '"Rp "* #,##0'

        ws.column_dimensions['A'].width = 7
        ws.column_dimensions['B'].width = 24
        for col_idx in range(3, len(col_numbers) + 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 22
        ws.row_dimensions[4].height = 42

    excel_buffer = io.BytesIO()
    col_numbers = ["1", "2", "3", "4", "5", "6", "7", "8", "9", "10", "11", "12", "13", "15"]

    with pd.ExcelWriter(excel_buffer, engine="openpyxl") as writer:
        df_tot_export = df_total_all.rename(columns=col_mapping)
        df_tot_export.to_excel(writer, sheet_name="UP3_TOTAL_KESELURUHAN", index=False, startrow=3)
        style_excel_sheet(writer.sheets["UP3_TOTAL_KESELURUHAN"], "REALISASI PEMUNGUTAN PPJ - TOTAL KESELURUHAN (UP3 MEDAN UTARA)", col_numbers)

        df_m_export = df_medan_raw.rename(columns=col_mapping)
        df_m_export.to_excel(writer, sheet_name="PEMDA_KOTA_MEDAN", index=False, startrow=3)
        style_excel_sheet(writer.sheets["PEMDA_KOTA_MEDAN"], "REALISASI PEMUNGUTAN PPJ - PEMDA KOTA MEDAN (UP3 MEDAN UTARA)", col_numbers)

        df_ds_export = df_deli_raw.rename(columns=col_mapping)
        df_ds_export.to_excel(writer, sheet_name="PEMDA_DELI_SERDANG", index=False, startrow=3)
        style_excel_sheet(writer.sheets["PEMDA_DELI_SERDANG"], "REALISASI PEMUNGUTAN PPJ - PEMDA KAB. DELI SERDANG (UP3 MEDAN UTARA)", col_numbers)

    st.download_button(
        label="📥 Download Laporan Excel Format Baru (UP3 Medan Utara)",
        data=excel_buffer.getvalue(),
        file_name="Laporan_Realisasi_PPJ_Format_Baru.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

else:
    st.info("💡 Unggah berkas Halaman 1 dan Halaman 2 format baru pada bilah kiri.")
